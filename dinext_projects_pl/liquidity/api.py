"""Server-authoritative report endpoints. Every call reloads permissions and source state."""

import json
from dataclasses import asdict
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import now, today

from .calculation import calculate
from .contracts import (
	CalendarFormInput,
	CalendarResponse,
	CellDetailsFormInput,
	CellDetailsResponse,
	ChildrenFormInput,
	ChildrenResponse,
	ExportFormInput,
	OptionResponse,
	OptionsFormInput,
	SettingsCatalogResponse,
	SettingsConflictResponse,
)
from .details import decorate_details
from .model import FilterInput, Filters, as_date
from .navigation import decorate_tree
from .presentation import Calendar, wire
from .sources import Loader, load, require_access


class CalendarSettingsChanged(frappe.ValidationError):
	http_status_code = 409


def require_version(snapshot, expected):
	if not snapshot.policy:
		return
	if expected is not None and (not isinstance(expected, str) or expected != snapshot.policy.version):
		raise CalendarSettingsChanged(_("Payment Calendar Settings changed. Refresh the calendar."))
	if expected is None and getattr(frappe.local, "request", None):
		raise CalendarSettingsChanged(_("Refresh the calendar before requesting this operation."))


def settings_conflict(error: CalendarSettingsChanged) -> SettingsConflictResponse:
	# A normal concurrency outcome has no traceback or Error Log. Other failures
	# still use Frappe's regular exception reporting and browser Console checks.
	if not getattr(frappe.local, "request", None):
		raise error
	frappe.response["http_status_code"] = 409
	return {"error": "calendar_settings_changed", "message": str(error)}


def calendar(
	filters: FilterInput | str | None,
	expected_settings_version=None,
	followup=False,
	expected_metrics=None,
	include_cash=True,
) -> Calendar:
	try:
		normalized = Filters.parse(filters, as_date(today()))
	except (ValueError, TypeError, OverflowError):
		frappe.throw(_("Invalid calendar filters"))
	loader, snapshot = load(normalized, include_cash=include_cash, expected_metrics=expected_metrics)
	if followup or expected_settings_version is not None:
		require_version(snapshot, expected_settings_version)
	flows = calculate(snapshot, include_zero=normalized.show_zero, expected_metrics=expected_metrics)
	# Restrict after allocation so hidden siblings cannot inflate the permitted share.
	visible = [f for f in flows if loader.allowed_analytic(f.analytic)]
	if len(visible) != len(flows):
		snapshot.can_balance = False
		snapshot.complete = False
		snapshot.warnings.append("Only permitted analytical shares are shown; balances are hidden")
	if snapshot.allocations:
		invoice_shares = {(i.source, s.key): s.analytic for i in snapshot.invoices for s in i.shares}
		request_shares = {(r.source, s.key): s.analytic for r in snapshot.requests for s in r.shares}
		snapshot.allocations = [
			a
			for a in snapshot.allocations
			if (a["invoice"], a["invoice_share"]) in invoice_shares
			and (a["request"], a["request_share"]) in request_shares
			and loader.allowed_analytic(invoice_shares[(a["invoice"], a["invoice_share"])])
			and loader.allowed_analytic(request_shares[(a["request"], a["request_share"])])
		]
	if not snapshot.complete:
		snapshot.can_balance = False
	return Calendar(normalized, snapshot, visible, _)


@frappe.whitelist()
def get_calendar(
	filters=None, depth=1, expected_settings_version=None
) -> CalendarResponse | SettingsConflictResponse:
	"""Calculate the authorized calendar and load all branches through the requested depth."""
	form_input: CalendarFormInput = {
		"filters": filters,
		"depth": depth,
		"expected_settings_version": expected_settings_version,
	}
	try:
		if isinstance(depth, bool) or str(depth) not in {str(i) for i in range(1, 8)}:
			raise ValueError
		depth = int(depth)
	except (ValueError, TypeError):
		frappe.throw(_("Invalid calendar depth"))
	try:
		calculated = calendar(form_input["filters"], expected_settings_version)
	except CalendarSettingsChanged as exc:
		return settings_conflict(exc)
	result = calculated.payload()
	result["rows"] = decorate_tree(calculated.tree(depth), calculated.filters)
	result["calculated_at"] = now()
	result["settings_version"] = calculated.snapshot.policy.version
	return result


@frappe.whitelist()
def get_children(
	filters, node_key="root", expected_settings_version=None
) -> ChildrenResponse | SettingsConflictResponse:
	"""Return the authorized children of a current calendar node."""
	form_input: ChildrenFormInput = {
		"filters": filters,
		"node_key": node_key,
		"expected_settings_version": expected_settings_version,
	}
	try:
		Calendar.validate_node_key(form_input["node_key"])
		calculated = calendar(form_input["filters"], expected_settings_version, followup=True)
		return {
			"rows": decorate_tree(calculated.children(form_input["node_key"]), calculated.filters),
			"settings_version": calculated.snapshot.policy.version,
		}
	except CalendarSettingsChanged as exc:
		return settings_conflict(exc)
	except ValueError as exc:
		frappe.throw(_(str(exc)))


@frappe.whitelist()
def get_cell_details(
	filters, node_key, period_key, metric, start=0, page_length=50, expected_settings_version=None
) -> CellDetailsResponse | SettingsConflictResponse:
	"""Return a bounded page of authorized components for one calendar cell."""
	form_input: CellDetailsFormInput = {
		"filters": filters,
		"node_key": node_key,
		"period_key": period_key,
		"metric": metric,
		"start": start,
		"page_length": page_length,
		"expected_settings_version": expected_settings_version,
	}
	try:
		Calendar.validate_cell_input(node_key, period_key, metric, start, page_length)
		expected_metrics = (
			("plan",)
			if metric == "plan"
			else ("forecast",)
			if metric in ("forecast", "combined", "projected_opening", "projected_closing")
			else ()
		)
		# A child row may exist only in another metric; retain its complete authorized
		# tree for strict node validation, even when the requested cell is zero.
		if node_key not in ("root", "net"):
			expected_metrics = None
		calculated = calendar(
			form_input["filters"],
			expected_settings_version,
			followup=True,
			expected_metrics=expected_metrics,
			include_cash=not (node_key in ("root", "net") and metric in ("plan", "forecast")),
		)
		result = calculated.details(
			form_input["node_key"],
			form_input["period_key"],
			form_input["metric"],
			form_input["start"],
			form_input["page_length"],
		)
		decorate_details(result, calculated.filters, calculated.snapshot)
		result["calculated_at"] = now()
		result["settings_version"] = calculated.snapshot.policy.version
		return result
	except CalendarSettingsChanged as exc:
		return settings_conflict(exc)
	except ValueError as exc:
		frappe.throw(_(str(exc)))


def safe_text(value):
	text = str(value or "")
	return (
		"'" + text
		if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n"))
		else text
	)


def make_workbook(result: Calendar) -> bytes:
	from openpyxl import Workbook
	from openpyxl.styles import Font

	wb = Workbook()
	ws = wb.active
	ws.title = _("Calendar")[:31]
	keys = (["undated"] if result.has_undated else []) + [p.key for p in result.periods] + ["total"]
	selected = result.filters.metrics
	if {"actual", "forecast"}.issubset(selected):
		cash_metric = "combined"
	else:
		cash_metric = next((metric for metric in ("actual", "forecast") if metric in selected), None)
	metrics = (["plan"] if "plan" in selected else []) + ([cash_metric] if cash_metric else [])
	columns = [(key, metric) for key in keys for metric in metrics]
	period_labels = {}
	for period in result.periods:
		if result.filters.grain == "week":
			year, week, _weekday = period.start.isocalendar()
			period_labels[period.key] = f"{week:02d} · {year}"
		elif result.filters.grain == "month":
			period_labels[period.key] = period.start.strftime("%m.%Y")
		else:
			period_labels[period.key] = period.start.strftime("%d.%m.%Y")
	period_labels.update(total=_("Total"), undated=_("Without date"))

	value = result.display_value

	metric_labels = {
		"combined": _("Cash movement"),
		"plan": _("Plan"),
		"forecast": _("Forecast"),
		"actual": _("Actual"),
		"transfer": _("Internal transfers"),
		"adjustment": _("Balance adjustments"),
	}

	def column_label(key, metric):
		if metric == "combined":
			kinds = (
				{result.period_metric(p.key) for p in result.periods}
				if key == "total"
				else {result.period_metric(key)}
			)
			if len(kinds) == 1:
				metric = kinds.pop()
		return metric_labels[metric]

	ws.append(
		[_("Name"), _("Level")]
		+ [f"{period_labels[key]} / {column_label(key, metric)}" for key, metric in columns]
	)
	for row in result.all_rows():
		ws.append(
			[safe_text(row["label"]), row["depth"]] + [value(row, key, metric) for key, metric in columns]
		)
		ws.row_dimensions[ws.max_row].outlineLevel = row["depth"]
		for cell in ws[ws.max_row][2:]:
			cell.number_format = "#,##0." + "0" * result.snapshot.precision
	ws.freeze_panes = "C2"
	ws.column_dimensions["A"].width = 48
	payload = result.payload()
	for label, values in ((_("Net cash movement"), payload["net"]),):
		ws.append([label, 0] + [value(values, k, m) for k, m in columns])
	if "actual" in selected:
		for metric in ("transfer", "adjustment"):
			if not any(f.metric == metric for f in result.visible):
				continue
			values = {}
			for period in result.periods:
				visible = cash_metric != "combined" or result.period_metric(period.key) == "actual"
				values[period.key] = payload["net"].get(period.key + "__" + metric, 0) if visible else 0
			values["total"] = sum(values.values())
			ws.append(
				[metric_labels[metric], 0]
				+ [values.get(key, 0) if col_metric == cash_metric else None for key, col_metric in columns]
			)

	detail = wb.create_sheet(_("Components")[:31])
	detail.append(
		[
			_(s)
			for s in (
				"Component",
				"Metric",
				"Direction",
				"Amount",
				"Currency",
				"Original amount",
				"Original due date",
				"Expected date",
				"Project",
				"Counterparty",
				"Contract",
				"Order",
				"Invoice",
				"Payment",
				"Source",
				"Status",
				"Paid",
				"Outstanding",
				"Links",
				"Date source",
				"Original due date overdue",
				"Match basis",
				"Covered amount",
				"Uncovered amount",
				"Amount source",
			)
		]
	)
	detail_metrics = list(metrics)
	if "actual" in selected:
		detail_metrics += [
			("combined_" if cash_metric == "combined" else "") + metric
			for metric in ("transfer", "adjustment")
		]
	for f in result.visible:
		if not any(result.matches_metric(f, metric) for metric in detail_metrics):
			continue
		a = f.analytic
		values = [
			f.id,
			f.metric,
			f.direction,
			float(f.amount),
			f.currency,
			float(f.original_amount),
			f.original_date.isoformat() if f.original_date else "",
			f.date.isoformat() if f.date else "",
			a.project,
			a.party,
			a.contract,
			a.order.name if a.order else "",
			a.invoice.name if a.invoice else "",
			f.payment.name if f.payment else "",
			f.source.name,
			f.status,
			float(f.paid),
			float(f.outstanding),
			json.dumps([asdict(link) for link in f.links], ensure_ascii=False),
			f"{f.date_source.doctype}: {f.date_source.name}" if f.date_source else f.source.name,
			int(f.contract_overdue),
			f.match_basis,
			float(f.covered),
			float(f.uncovered),
			f"{f.amount_source.doctype}: {f.amount_source.name}" if f.amount_source else f.source.name,
		]
		detail.append([safe_text(v) if isinstance(v, str) else v for v in values])
	detail.freeze_panes = "D2"
	parameters = wb.create_sheet(_("Parameters")[:31])
	parameters.append([_("Calculation date"), result.snapshot.today.isoformat()])
	parameters.append([_("Company currency"), result.snapshot.currency])
	if result.snapshot.policy:
		parameters.append([_("Settings version"), result.snapshot.policy.version])
		parameters.append(
			[
				_("Calculation method"),
				safe_text(json.dumps(result.snapshot.policy.as_dict(), ensure_ascii=False)),
			]
		)
		allocations = wb.create_sheet(_("Allocations")[:31])
		allocations.append(
			[
				_(v)
				for v in (
					"Metric",
					"Payment Request",
					"Invoice",
					"Match basis",
					"Currency",
					"Covered amount",
					"Remaining amount",
				)
			]
		)
		for row in result.snapshot.allocations:
			allocations.append(
				[
					safe_text(row["metric"]),
					safe_text(row["request"].name),
					safe_text(row["invoice"].name),
					safe_text(row["basis"]),
					safe_text(row["currency"]),
					float(row["covered_original"]),
					float(row["remaining_original"]),
				]
			)
	for key, value in wire(asdict(result.filters)).items():
		parameters.append([key, safe_text(json.dumps(value, ensure_ascii=False))])
	for warning in payload["warnings"]:
		parameters.append([_("Warning"), safe_text(warning)])
	if payload["balances"]:
		balance = wb.create_sheet(_("Balances")[:31])
		balance.append(
			[
				_(s)
				for s in (
					"Period",
					"Actual opening",
					"Actual closing",
					"Projected opening",
					"Projected closing",
				)
			]
		)
		for row in payload["balances"]["periods"]:
			balance.append(list(row.values()))
	for sheet in wb:
		for cell in sheet[1]:
			cell.font = Font(bold=True)
	stream = BytesIO()
	wb.save(stream)
	return stream.getvalue()


@frappe.whitelist()
def export_liquidity_calendar(filters, expected_settings_version=None) -> None | SettingsConflictResponse:
	"""Download all authorized calendar rows and components as an XLSX attachment."""
	form_input: ExportFormInput = {"filters": filters, "expected_settings_version": expected_settings_version}
	try:
		result = calendar(form_input["filters"], expected_settings_version, followup=True)
	except CalendarSettingsChanged as exc:
		return settings_conflict(exc)
	frappe.response.filename = "liquidity-payment-calendar.xlsx"
	frappe.response.filecontent = make_workbook(result)
	frappe.response.type = "binary"
	frappe.response.display_content_as = "attachment"


@frappe.whitelist()
def get_options(kind, company=None, txt="", party_type=None) -> list[OptionResponse]:
	"""Suggest permitted calendar filter values within the selected company."""
	form_input: OptionsFormInput = {
		"kind": kind,
		"company": company,
		"txt": txt,
		"party_type": party_type,
	}
	if (
		not isinstance(kind, str)
		or kind not in ("company", "project", "portfolio", "account", "party")
		or (
			company is not None
			and (
				not isinstance(company, str)
				or len(company) > 200
				or (kind != "company" and not company.strip())
			)
		)
		or (
			party_type is not None
			and (
				not isinstance(party_type, str) or party_type not in ("", "Customer", "Supplier", "Employee")
			)
		)
		or not isinstance(txt, str)
		or len(txt) > 200
	):
		frappe.throw(_("Invalid calendar selection"))
	kind, company, txt, party_type = (
		form_input["kind"],
		form_input["company"],
		form_input["txt"],
		form_input["party_type"],
	)
	if kind == "company":
		roles = set(frappe.get_roles())
		from .sources import FINANCE_ROLES, PROJECT_ROLES

		if not roles.intersection(FINANCE_ROLES | PROJECT_ROLES):
			frappe.throw(_("Not permitted"), frappe.PermissionError)
		doctype, conditions = "Company", {}
	else:
		full_role = require_access(company)
		if kind == "project":
			doctype, conditions = "Project", {"company": company}
		elif kind == "account":
			doctype, conditions = (
				"Account",
				{"company": company, "is_group": 0, "account_type": ["in", ["Bank", "Cash"]]},
			)
		elif kind == "portfolio":
			loader = Loader(Filters.parse({"company": company}, as_date(today())))
			projects = loader.parents("Project", ["custom_project_portfolio"], {"company": company})
			values = sorted(
				{
					p.custom_project_portfolio
					for p in projects
					if p.custom_project_portfolio and txt.casefold() in p.custom_project_portfolio.casefold()
				}
			)
			return [{"value": v, "description": ""} for v in values[:50]]
		else:
			if party_type not in ("Customer", "Supplier", "Employee"):
				return []
			# Suggestions use the same permitted analytical shares as the report, not
			# the global party directory (which can include other companies/projects).
			result = calendar({"company": company})
			names = sorted(
				{
					f.analytic.party
					for f in result.flows
					if f.analytic.party_type == party_type
					and f.analytic.party
					and txt.casefold() in f.analytic.party.casefold()
				}
			)
			return [{"value": name, "description": ""} for name in names[:50]]
		if kind == "account" and not full_role:
			result = calendar({"company": company})
			used = {f.account for f in result.flows if f.account}
			if not used:
				return []
			conditions["name"] = ["in", sorted(used)]
	if not frappe.has_permission(doctype, "read"):
		return []
	rows = frappe.get_list(
		doctype,
		filters=conditions,
		or_filters={"name": ["like", "%" + txt + "%"]},
		fields=["name"],
		order_by="name",
		limit_page_length=50,
	)
	return [
		{"value": r.name, "description": ""} for r in rows if frappe.has_permission(doctype, "read", r.name)
	]


@frappe.whitelist()
def get_settings_catalog() -> SettingsCatalogResponse:
	"""Return supported source choices to authorized calendar configuration managers."""
	from .settings import MANAGERS, SETTINGS, SUPPORTED, date_fields, status_options

	if not MANAGERS.intersection(frappe.get_roles()) or not frappe.has_permission(SETTINGS, "write"):
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	return {
		"sources": {
			dt: {
				"dates": [
					{
						"value": name,
						"label": _("Payment Schedule")
						if name == "payment_schedule"
						else _(frappe.get_meta(dt).get_label(name) or name),
					}
					for name in date_fields(dt)
				],
				"statuses": list(status_options(dt)),
			}
			for dt in SUPPORTED
		}
	}
