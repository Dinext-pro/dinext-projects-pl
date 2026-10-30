# Copyright (c) 2026, Dinext and contributors
# For license information, please see license.txt

from dataclasses import asdict
from io import BytesIO

import frappe
from frappe.tests.utils import FrappeTestCase
from frappe.utils import add_days, today
from openpyxl import load_workbook

from dinext_projects_pl.liquidity.api import (
	CalendarSettingsChanged,
	calendar,
	get_calendar,
	get_cell_details,
	get_children,
	get_settings_catalog,
	make_workbook,
)
from dinext_projects_pl.liquidity.settings import (
	SETTINGS,
	compile_policy,
	initial_profile,
	install_settings,
	load_policy,
)
from dinext_projects_pl.tests import liquidity_fixtures as fixtures


class TestPaymentCalendarSettings(FrappeTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		frappe.set_user("Administrator")
		cls.context = fixtures.setup_company()
		install_settings()

	def setUp(self):
		frappe.set_user("Administrator")
		frappe.flags.mute_emails = True
		frappe.db.savepoint("calendar_settings_case")
		self.addCleanup(lambda: frappe.db.rollback(save_point="calendar_settings_case"))
		self.filters = {"company": fixtures.COMPANY, "from_date": today(), "to_date": add_days(today(), 90)}
		self.settings = frappe.get_doc(SETTINGS)
		self.settings.update(initial_profile())
		self.settings.save()

	def source(self, doctype, metric="Forecast"):
		return next(r for r in self.settings.sources if r.source_doctype == doctype and r.metric == metric)

	def test_metadata_has_only_two_tables_and_manager_permissions(self):
		meta = frappe.get_meta(SETTINGS)
		self.assertTrue(meta.issingle)
		self.assertTrue(meta.track_changes)
		self.assertEqual({f.fieldname for f in meta.get_table_fields()}, {"sources", "links"})
		self.assertEqual(
			{p.role for p in meta.permissions if p.write}, {"System Manager", "Accounts Manager"}
		)
		self.assertFalse(any(f.fieldname.startswith("fifo") for f in meta.fields))

	def test_saved_source_states_dates_and_links_survive_seed_and_reload(self):
		row = self.source("Payment Entry")
		row.document_state = "Both"
		row.date_field = "reference_date"
		row.fallback_date_field = "posting_date"
		row.statuses = "Draft\nSubmitted"
		self.settings.save()
		version = load_policy().version
		install_settings()
		self.assertEqual(load_policy().version, version)
		self.settings.reload()
		self.assertEqual(self.source("Payment Entry").date_field, "reference_date")
		self.assertEqual(
			compile_policy(self.settings).source("forecast", "Payment Entry").docstatuses, (0, 1)
		)

	def test_restore_after_replacing_child_rows_preserves_saved_rules(self):
		from dinext_projects_pl.tests.liquidity_ui_helpers import settings_values

		row = self.source("Payment Entry")
		row.date_field = "reference_date"
		row.document_state = "Both"
		self.settings.save()
		expected = load_policy().version
		backup = settings_values(self.settings)
		self.settings.update(initial_profile())
		self.settings.save()
		self.assertNotEqual(load_policy().version, expected)
		self.settings.update(backup)
		self.settings.save()
		self.settings.reload()
		self.assertEqual(len(self.settings.sources), 14)
		self.assertEqual(load_policy().version, expected)

	def test_omitted_sources_are_disabled_in_a_valid_sparse_profile(self):
		order = fixtures.sales_order(self.context)
		invoice = fixtures.sales_invoice(order, 60)
		fixtures.pay(self.context, invoice, 20)
		baseline = get_calendar(self.filters)
		row = self.source("Sales Invoice").as_dict()
		self.settings.set("sources", [row])
		self.settings.set("links", [])
		self.settings.save()
		self.assertIsNone(load_policy().source("plan", "Sales Order"))
		result = get_calendar(self.filters)
		self.assertEqual(result["directions"]["in"].get("total__plan", 0), 0)
		self.assertEqual(result["directions"]["in"]["total__forecast"], 40)
		self.assertEqual(
			result["directions"]["in"]["total__actual"],
			baseline["directions"]["in"]["total__actual"],
		)
		children = get_children(self.filters, expected_settings_version=result["settings_version"])
		self.assertEqual(children["settings_version"], result["settings_version"])
		details = get_cell_details(
			self.filters,
			"root",
			"total",
			"forecast",
			expected_settings_version=result["settings_version"],
		)
		self.assertEqual(details["total"], 40)
		self.assertEqual(details["settings_version"], result["settings_version"])
		workbook = load_workbook(BytesIO(make_workbook(calendar(self.filters))))
		self.assertIn("Allocations", workbook.sheetnames)

	def test_invalid_metadata_duplicate_state_and_cycle_are_rejected(self):
		for field, value in (
			("date_field", "grand_total"),
			("fallback_date_field", "not_a_date"),
			("statuses", "Unknown"),
		):
			with self.subTest(field=field):
				doc = frappe.get_doc(SETTINGS)
				next(
					r for r in doc.sources if r.source_doctype == "Sales Invoice" and r.metric == "Forecast"
				).set(field, value)
				with self.assertRaises(frappe.ValidationError):
					doc.save()
		self.settings.append(
			"links",
			{
				"metric": "Forecast",
				"source_doctype": "Payment Entry",
				"target_doctype": "Sales Invoice",
				"enabled": 1,
			},
		)
		with self.assertRaises(frappe.ValidationError):
			self.settings.save()

	def test_real_fifo_paid_invoice_and_export(self):
		order = fixtures.sales_order(self.context)
		request = fixtures.payment_request(order, 100, add_days(today(), 14))
		invoice = fixtures.sales_invoice(order, 60)
		fixtures.pay(self.context, invoice, 20)
		result = get_calendar(self.filters)
		self.assertEqual(result["directions"]["in"]["total__forecast"], 80)
		self.assertEqual(result["directions"]["in"]["total__actual"], 20)
		details = get_cell_details(
			self.filters, "root", "total", "forecast", expected_settings_version=result["settings_version"]
		)
		self.assertEqual(details["total"], 80)
		self.assertTrue(any(r["match_basis"] == "FIFO" for r in details["rows"]))
		book = load_workbook(BytesIO(make_workbook(calendar(self.filters))))
		self.assertTrue(any(request.name in [c.value for c in row] for sheet in book for row in sheet))
		fixtures.pay(self.context, invoice, 40)
		self.assertEqual(get_calendar(self.filters)["directions"]["in"]["total__forecast"], 40)

	def test_draft_entry_refines_then_submission_and_cancel_reconcile(self):
		from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry

		order = fixtures.sales_order(self.context)
		invoice = fixtures.sales_invoice(order, 100)
		entry = get_payment_entry(
			invoice.doctype, invoice.name, bank_account=self.context["bank"], party_amount=40
		)
		entry.reference_no = "Calendar settings regression"
		entry.reference_date = today()
		entry.posting_date = add_days(today(), 14)
		entry.insert()
		rows = get_cell_details(self.filters, "root", "total", "forecast")["rows"]
		refined = [r for r in rows if r["date_source"] == {"doctype": entry.doctype, "name": entry.name}]
		self.assertEqual(sum(r["amount"] for r in refined), 40)
		self.assertTrue(any(p["docstatus"] == 0 for r in refined for p in r["payment_stages"]))
		entry.posting_date = today()
		entry.save()
		entry.submit()
		self.assertEqual(get_calendar(self.filters)["directions"]["in"]["total__forecast"], 60)
		self.assertEqual(get_calendar(self.filters)["directions"]["in"]["total__actual"], 40)
		entry.cancel()
		self.assertEqual(get_calendar(self.filters)["directions"]["in"]["total__forecast"], 100)

	def test_source_disable_preserves_actual_cash_and_changes_projected_opening(self):
		order = fixtures.sales_order(self.context)
		invoice = fixtures.sales_invoice(order, 100)
		fixtures.pay(self.context, invoice, 20)
		filters = dict(self.filters, from_date=add_days(today(), 20))
		before = get_calendar(filters)
		self.assertEqual(before["balances"]["periods"][0]["projected_opening"], 100)
		self.source("Sales Invoice").enabled = 0
		self.settings.save()
		after = get_calendar(filters)
		self.assertEqual(after["balances"]["periods"][0]["projected_opening"], 20)
		self.assertEqual(after["balances"]["current_cash"], before["balances"]["current_cash"])
		self.assertEqual(after["balances"]["current_cash"], 20)

	def test_assign_invoice_to_both_metrics_and_change_date(self):
		order = fixtures.sales_order(self.context)
		fixtures.sales_invoice(order, 60)
		self.source("Sales Order", "Plan").enabled = 0
		self.source("Sales Invoice", "Plan").enabled = 1
		self.source("Sales Invoice", "Plan").date_field = "posting_date"
		self.settings.save()
		result = get_calendar(self.filters)
		self.assertEqual(result["directions"]["in"]["total__plan"], 60)
		self.assertEqual(result["directions"]["in"]["total__forecast"], 60)
		rows = get_cell_details(self.filters, "root", "total", "plan")["rows"]
		self.assertTrue(all(r["date"] == today() for r in rows))

	def test_payment_order_uses_fifo_with_disabled_request_source(self):
		order = fixtures.sales_order(self.context)
		request = fixtures.payment_request(order, 100, add_days(today(), 14))
		invoice = fixtures.sales_invoice(order, 60)
		fixtures.pay(self.context, invoice, 20)
		po = fixtures.payment_order(self.context, request)
		po.db_set("posting_date", add_days(today(), 25))
		self.source("Payment Request").enabled = 0
		self.settings.save()
		rows = get_cell_details(self.filters, "root", "total", "forecast")["rows"]
		self.assertEqual(sum(r["amount"] for r in rows), 80)
		self.assertTrue(all(r["date_source"] == {"doctype": po.doctype, "name": po.name} for r in rows))
		self.assertEqual(sum(r["amount"] for r in rows if r["analytic"]["invoice"]), 40)
		self.source("Sales Invoice").enabled = 0
		self.settings.save()
		rows = get_cell_details(self.filters, "root", "total", "forecast")["rows"]
		self.assertEqual(sum(r["amount"] for r in rows), 80)
		self.assertEqual(sum(r["amount"] for r in rows if r["analytic"]["invoice"]), 40)

	def test_partial_order_does_not_repeat_invoice_settlements(self):
		order = fixtures.sales_order(self.context)
		invoice = fixtures.sales_invoice(order, 60)
		fixtures.pay(self.context, invoice, 20)
		request = fixtures.payment_request(invoice, 40, add_days(today(), 14))
		po = fixtures.payment_order(self.context, request)
		po.db_set("posting_date", add_days(today(), 25))
		rows = get_cell_details(self.filters, "root", "total", "forecast")["rows"]
		self.assertEqual(sum(r["amount"] for r in rows), 40)
		self.assertTrue(all(r["date_source"] == {"doctype": po.doctype, "name": po.name} for r in rows))

	def test_stale_version_is_rejected_for_all_followups(self):
		order = fixtures.sales_order(self.context)
		fixtures.sales_invoice(order)
		version = get_calendar(self.filters)["settings_version"]
		self.source("Sales Invoice").date_field = "posting_date"
		self.settings.save()
		for operation in (
			lambda: get_children(self.filters, expected_settings_version=version),
			lambda: get_cell_details(
				self.filters, "root", "total", "forecast", expected_settings_version=version
			),
			lambda: calendar(self.filters, version, followup=True),
		):
			with self.assertRaises(CalendarSettingsChanged):
				operation()
		self.assertNotEqual(get_calendar(self.filters)["settings_version"], version)

	def test_catalog_restricted_and_all_dates_come_from_metadata(self):
		catalog = get_settings_catalog()
		self.assertIn("reference_date", {d["value"] for d in catalog["sources"]["Payment Entry"]["dates"]})
		previous = frappe.session.user
		try:
			frappe.set_user("Guest")
			self.assertFalse(frappe.has_permission(SETTINGS, "read"))
			with self.assertRaises(frappe.PermissionError):
				get_settings_catalog()
		finally:
			frappe.set_user(previous)
