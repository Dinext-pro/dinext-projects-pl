import json
from io import BytesIO
from unittest.mock import patch

import frappe
from frappe.tests.utils import FrappeTestCase
from frappe.utils import add_days, today
from openpyxl import load_workbook

from dinext_projects_pl.liquidity.api import calendar, get_calendar, get_cell_details, make_workbook
from dinext_projects_pl.tests import liquidity_fixtures as fixtures


def deny_test_project(doc, **kwargs):
	# A policy can inspect a field not requested by the report's parent query.
	return False if doc.get("project_name") == fixtures.PROJECT else None


class TestLiquidityIntegration(FrappeTestCase):
	def test_tree_navigation_targets_metadata_and_lazy_children(self):
		from dinext_projects_pl.liquidity.api import get_children

		order = fixtures.sales_order(self.context)
		contract = fixtures.contract(order)
		invoice = fixtures.sales_invoice(order)
		payment = fixtures.pay(self.context, invoice, 20)
		expected = {
			("Project", self.context["project"]),
			("Customer", order.customer),
			("Contract", contract.name),
			(order.doctype, order.name),
			(invoice.doctype, invoice.name),
			(payment.doctype, payment.name),
		}

		def flatten(rows):
			return [node for row in rows for node in [row, *flatten(row.get("children", []))]]

		for grouping in ("project_party", "party_project", "party_only"):
			filters = dict(self.filters, grouping=grouping)
			rows = flatten(get_calendar(filters, depth=7)["rows"])
			links = {(r["document"]["doctype"], r["document"]["name"]) for r in rows if r.get("document")}
			self.assertEqual(
				links,
				expected - ({("Project", self.context["project"])} if grouping == "party_only" else set()),
			)
			for row in rows:
				if row.get("document"):
					self.assertEqual(
						row["icon"], frappe.get_meta(row["document"]["doctype"]).get("icon") or ""
					)
				else:
					self.assertEqual(row["level"], "direction")
			root = get_calendar(filters)["rows"][0]
			lazy = get_children(filters, root["id"])
			self.assertTrue(lazy[0].get("document"))

	def test_tree_navigation_does_not_link_unreadable_counterparty(self):
		order = fixtures.sales_order(self.context)
		fixtures.sales_invoice(order)
		fixtures.sales_order(self.context, project=False)
		baseline = get_calendar(self.filters, depth=3)
		permission = frappe.has_permission
		with patch(
			"frappe.has_permission",
			side_effect=lambda dt, *a, **kw: False if dt == "Customer" else permission(dt, *a, **kw),
		):
			result = get_calendar(self.filters, depth=3)
		self.assertEqual(result["directions"], baseline["directions"])
		projects = result["rows"][0]["children"]
		self.assertTrue(any(not row["path"][-1] and "document" not in row for row in projects))
		for project in projects:
			party = project["children"][0]
			self.assertEqual(party["label"], order.customer)
			self.assertNotIn("document", party)
			self.assertNotIn("icon", party)

	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		frappe.set_user("Administrator")
		cls.context = fixtures.setup_company()
		cls.manager = "_test_liquidity_project@example.invalid"
		if not frappe.db.exists("User", cls.manager):
			frappe.get_doc(
				{
					"doctype": "User",
					"email": cls.manager,
					"first_name": "Liquidity Project Test",
					"send_welcome_email": 0,
					"roles": [
						{"role": role}
						for role in ("Dinext_Project_Manager", "Projects Manager", "Dinext_User")
						if frappe.db.exists("Role", role)
					],
				}
			).insert()
		from frappe.share import add

		for dt, name in (("Company", fixtures.COMPANY), ("Project", cls.context["project"])):
			add(dt, name, user=cls.manager, read=1)
		frappe.get_doc(
			{
				"doctype": "User Permission",
				"user": cls.manager,
				"allow": "Project",
				"for_value": cls.context["project"],
				"apply_to_all_doctypes": 1,
			}
		).insert()

	def setUp(self):
		super().setUp()
		frappe.set_user("Administrator")
		frappe.db.savepoint("liquidity_case")
		self.addCleanup(lambda: frappe.db.rollback(save_point="liquidity_case"))
		self.filters = {"company": fixtures.COMPANY, "from_date": today(), "to_date": add_days(today(), 90)}

	def totals(self):
		return get_calendar(self.filters)["directions"]["in"]

	def test_depth_api_calculates_once_and_rejects_invalid_depth_before_loading(self):
		order = fixtures.sales_order(self.context)
		fixtures.sales_invoice(order)
		with patch("dinext_projects_pl.liquidity.api.calendar", wraps=calendar) as calculate_calendar:
			result = get_calendar(dict(self.filters, grouping="party_only"), depth="6")
			calculate_calendar.assert_called_once()
		self.assertEqual(result["rows"][0]["children"][0]["level"], "party")
		for depth in (0, 8, True, None, [], "2.5", "2.0"):
			with (
				self.subTest(depth=depth),
				patch("dinext_projects_pl.liquidity.api.calendar") as calculate_calendar,
			):
				with self.assertRaises(frappe.ValidationError):
					get_calendar(self.filters, depth=depth)
				calculate_calendar.assert_not_called()

	def test_party_only_batch_respects_denied_project(self):
		order = fixtures.sales_order(self.context)
		fixtures.sales_invoice(order)
		frappe.set_user(self.manager)
		try:
			ordinary = get_calendar(self.filters, depth=7)
			party_only = get_calendar(dict(self.filters, grouping="party_only"), depth=6)
			self.assertIsNone(party_only["balances"])
			self.assertEqual(party_only["directions"], ordinary["directions"])
			self.assertNotIn('"level": "project"', json.dumps(party_only["rows"]))
		finally:
			frappe.set_user("Administrator")

	def test_real_order_invoice_payment_and_contract(self):
		order = fixtures.sales_order(self.context)
		contract = fixtures.contract(order)
		invoice = fixtures.sales_invoice(order)
		payment = fixtures.pay(self.context, invoice, 20)
		result = calendar(self.filters)
		payload = result.payload()
		self.assertEqual(payload["directions"]["in"]["total__plan"], 80)
		self.assertEqual(payload["directions"]["in"]["total__forecast"], 40)
		self.assertEqual(payload["directions"]["in"]["total__actual"], 20)
		self.assertEqual(result.details("root", "total", "actual")["rows"][0]["paid"], 20)
		self.assertEqual(payload["balances"]["current_cash"], 20)
		self.assertEqual(payload["balances"]["periods"][-1]["projected_closing"], 60)
		self.assertTrue(any(r["label"] == contract.name for r in result.all_rows()))
		self.assertTrue(any(r["label"] == payment.name for r in result.all_rows()))
		workbook = load_workbook(BytesIO(make_workbook(result)))
		self.assertGreater(len(workbook.worksheets[1]["A"]), 1)
		self.assertEqual(workbook.worksheets[1]["D2"].data_type, "n")
		from dinext_projects_pl.dinext_projects_pl.report.liquidity_payment_calendar.liquidity_payment_calendar import (
			execute,
		)

		columns, rows, _, chart, summary, _ = execute(self.filters)
		self.assertTrue(any(c["fieldname"] == "total__combined" for c in columns))
		self.assertEqual(rows[0]["total__combined"], 40)
		self.assertEqual([sum(series["values"]) for series in chart["data"]["datasets"]], [80, 0, 40, 0])
		self.assertEqual(summary[0]["value"], 40)

	def test_real_request_without_invoice_and_partial_payment(self):
		order = fixtures.sales_order(self.context)
		request = fixtures.payment_request(order, 100, add_days(today(), 14))
		self.assertEqual(self.totals()["total__forecast"], 100)
		fixtures.pay(self.context, order, 30, request)
		result = get_calendar(self.filters)
		self.assertEqual(result["directions"]["in"]["total__forecast"], 70)
		self.assertEqual(result["directions"]["in"]["total__actual"], 30)
		self.assertEqual(result["directions"]["in"]["total__plan"], 70)

	def test_real_invoice_request_refines_date(self):
		order = fixtures.sales_order(self.context)
		invoice = fixtures.sales_invoice(order, 100)
		request = fixtures.payment_request(invoice, 60, add_days(today(), 21))
		fixtures.pay(self.context, invoice, 20, request)
		rows = get_cell_details(self.filters, "root", "total", "forecast")["rows"]
		self.assertEqual(sum(r["amount"] for r in rows), 80)
		self.assertEqual(sum(r["amount"] for r in rows if r["date"] == add_days(today(), 21)), 40)
		self.assertEqual(sum(r["amount"] for r in rows if r["date"] == add_days(today(), 7)), 40)

	def test_cancelled_payment_restores_forecast(self):
		order = fixtures.sales_order(self.context)
		invoice = fixtures.sales_invoice(order, 100)
		payment = fixtures.pay(self.context, invoice, 20)
		payment.cancel()
		result = self.totals()
		self.assertEqual(result["total__forecast"], 100)
		self.assertEqual(result.get("total__actual", 0), 0)

	def test_closed_order_keeps_only_billed_unpaid_plan(self):
		order = fixtures.sales_order(self.context)
		invoice = fixtures.sales_invoice(order, 60)
		fixtures.pay(self.context, invoice, 20)
		order.reload()
		order.update_status("Closed")
		result = self.totals()
		self.assertEqual(result["total__plan"], 40)
		self.assertEqual(result["total__forecast"], 40)
		self.assertEqual(result["total__actual"], 20)

	def test_cancelled_and_draft_requests_do_not_refine_invoice(self):
		order = fixtures.sales_order(self.context)
		invoice = fixtures.sales_invoice(order, 60)
		request = fixtures.payment_request(invoice, 60, add_days(today(), 40))
		request.cancel()
		draft = frappe.copy_doc(request)
		draft.docstatus = 0
		draft.status = "Draft"
		draft.insert()
		rows = get_cell_details(self.filters, "root", "total", "forecast")["rows"]
		self.assertEqual(sum(r["amount"] for r in rows), 60)
		self.assertTrue(all(r["date"] == str(invoice.due_date) for r in rows))
		self.assertNotIn(request.name, str(rows))
		self.assertNotIn(draft.name, str(rows))

	def test_workspace_install_preserves_optional_links_and_custom_content(self):
		import json

		from dinext_projects_pl.liquidity.setup import REPORT, install_report, install_workspace_links

		install_report()
		content = [{"id": "keep", "type": "header", "data": {"text": "Keep custom content", "col": 12}}]
		workspace = frappe.get_doc(
			{
				"doctype": "Workspace",
				"label": "_Test Liquidity Workspace",
				"title": "_Test Liquidity Workspace",
				"module": "Dinext Projects Pl",
				"public": 1,
				"content": json.dumps(content),
				"number_cards": [{"label": "Optional", "number_card_name": "_Test Absent Optional Card"}],
			}
		)
		old = frappe.conf.developer_mode
		try:
			frappe.conf.developer_mode = 0
			workspace.flags.ignore_links = True
			workspace.insert()
			install_workspace_links([workspace.name])
			workspace.reload()
			self.assertEqual(workspace.number_cards[0].number_card_name, "_Test Absent Optional Card")
			self.assertEqual(json.loads(workspace.content)[0], content[0])
			self.assertEqual(sum(r.link_to == REPORT for r in workspace.links), 1)
			first = workspace.content
			install_workspace_links([workspace.name])
			workspace.reload()
			self.assertEqual(workspace.content, first)
		finally:
			frappe.conf.developer_mode = old

	def test_direct_guest_access_denied(self):
		frappe.set_user("Guest")
		try:
			with self.assertRaises(frappe.PermissionError):
				get_calendar(dict(self.filters, grouping="party_only"), depth=7)
		finally:
			frappe.set_user("Administrator")

	def test_advance_consumed_by_invoice_once(self):
		order = fixtures.sales_order(self.context)
		fixtures.pay(self.context, order, 20)
		fixtures.sales_invoice(order, 60, advances=True)
		result = self.totals()
		self.assertEqual(result["total__plan"], 80)
		self.assertEqual(result["total__forecast"], 40)
		self.assertEqual(result["total__actual"], 20)

	def test_multiple_invoices_and_payment_allocation(self):
		order = fixtures.sales_order(self.context)
		first = fixtures.sales_invoice(order, 60)
		second = fixtures.sales_invoice(order, 40)
		fixtures.pay(self.context, first, 20)
		fixtures.pay(self.context, second, 10)
		result = self.totals()
		self.assertEqual(result["total__plan"], 70)
		self.assertEqual(result["total__forecast"], 70)
		self.assertEqual(result["total__actual"], 30)
		rows = get_cell_details(self.filters, "root", "total", "actual")["rows"]
		self.assertEqual(
			{r["analytic"]["invoice"]["name"]: r["amount"] for r in rows}, {first.name: 20, second.name: 10}
		)

	def test_internal_transfer_and_fee_reconcile_selected_accounts(self):
		expense = frappe.db.get_value("Company", fixtures.COMPANY, "default_expense_account")
		fixtures.journal(
			self.context,
			[
				{"account": self.context["bank"], "credit_in_account_currency": 103},
				{"account": self.context["bank_b"], "debit_in_account_currency": 100},
				{
					"account": expense,
					"debit_in_account_currency": 3,
					"cost_center": frappe.db.get_value("Company", fixtures.COMPANY, "cost_center"),
				},
			],
		)
		result = get_calendar(self.filters)
		self.assertEqual(result["directions"]["in"].get("total__actual", 0), 0)
		self.assertEqual(result["directions"]["out"]["total__actual"], 3)
		self.assertEqual(result["balances"]["current_cash"], -3)
		result = get_calendar(dict(self.filters, accounts=[self.context["bank"]]))
		self.assertEqual(result["directions"]["out"]["total__transfer"], 100)
		self.assertEqual(result["balances"]["current_cash"], -103)
		self.assertIsNone(result["balances"]["minimum"])

	def test_payment_order_journal_execution(self):
		invoice = fixtures.purchase_invoice(self.context)
		request = fixtures.payment_request(invoice, 100, add_days(today(), 14))
		order = fixtures.payment_order(self.context, request)
		before = get_calendar(self.filters)
		self.assertEqual(before["directions"]["out"]["total__forecast"], 100)
		self.assertEqual(before["directions"]["out"].get("total__actual", 0), 0)
		fixtures.journal(
			self.context,
			[
				{
					"account": invoice.credit_to,
					"party_type": "Supplier",
					"party": fixtures.SUPPLIER,
					"debit_in_account_currency": 30,
					"reference_type": "Purchase Invoice",
					"reference_name": invoice.name,
				},
				{"account": self.context["bank"], "credit_in_account_currency": 30},
			],
			payment_order=order.name,
		)
		after = get_calendar(self.filters)
		self.assertEqual(after["directions"]["out"]["total__forecast"], 70)
		self.assertEqual(after["directions"]["out"]["total__actual"], 30)
		rows = get_cell_details(self.filters, "root", "total", "forecast")["rows"]
		self.assertTrue(any(link["name"] == order.name for row in rows for link in row["links"]))

	def test_internal_transfer_mixed_with_supplier_payment(self):
		invoice = fixtures.purchase_invoice(self.context)
		fixtures.journal(
			self.context,
			[
				{"account": self.context["bank"], "credit_in_account_currency": 130},
				{"account": self.context["bank_b"], "debit_in_account_currency": 100},
				{
					"account": invoice.credit_to,
					"party_type": "Supplier",
					"party": fixtures.SUPPLIER,
					"debit_in_account_currency": 30,
					"reference_type": "Purchase Invoice",
					"reference_name": invoice.name,
				},
			],
		)
		result = get_calendar(self.filters)
		self.assertEqual(result["directions"]["in"].get("total__actual", 0), 0)
		self.assertEqual(result["directions"]["out"]["total__actual"], 30)
		self.assertEqual(result["directions"]["out"]["total__forecast"], 70)
		self.assertEqual(result["balances"]["current_cash"], -30)

	def test_non_cash_settlement_does_not_create_actual(self):
		order = fixtures.sales_order(self.context)
		invoice = fixtures.sales_invoice(order, 100)
		expense = frappe.db.get_value("Company", fixtures.COMPANY, "default_expense_account")
		fixtures.journal(
			self.context,
			[
				{
					"account": invoice.debit_to,
					"party_type": "Customer",
					"party": fixtures.CUSTOMER,
					"credit_in_account_currency": 30,
					"reference_type": "Sales Invoice",
					"reference_name": invoice.name,
				},
				{
					"account": expense,
					"debit_in_account_currency": 30,
					"cost_center": frappe.db.get_value("Company", fixtures.COMPANY, "cost_center"),
				},
			],
			voucher_type="Write Off Entry",
		)
		result = self.totals()
		self.assertEqual(result["total__forecast"], 70)
		self.assertEqual(result["total__plan"], 70)
		self.assertEqual(result.get("total__actual", 0), 0)

	def test_return_credit_reduces_debt_without_cash(self):
		from erpnext.accounts.doctype.sales_invoice.sales_invoice import make_sales_return

		order = fixtures.sales_order(self.context)
		invoice = fixtures.sales_invoice(order, 100)
		credit = make_sales_return(invoice.name)
		credit.update_outstanding_for_self = 0
		credit.items[0].qty = -0.3
		credit.set("payment_schedule", [])
		credit.insert()
		credit.submit()
		result = self.totals()
		self.assertEqual(result["total__forecast"], 70)
		self.assertEqual(result.get("total__actual", 0), 0)
		credit.cancel()
		self.assertEqual(self.totals()["total__forecast"], 100)

	def test_credit_note_with_own_debt_forecasts_refund(self):
		from erpnext.accounts.doctype.sales_invoice.sales_invoice import make_sales_return

		order = fixtures.sales_order(self.context)
		invoice = fixtures.sales_invoice(order, 100)
		credit = make_sales_return(invoice.name)
		credit.update_outstanding_for_self = 1
		credit.items[0].qty = -0.3
		credit.set("payment_schedule", [])
		credit.insert()
		credit.submit()
		result = get_calendar(self.filters)
		self.assertEqual(result["directions"]["in"]["total__forecast"], 100)
		self.assertEqual(result["directions"]["out"]["total__forecast"], 30)
		self.assertEqual(result["directions"]["out"].get("total__actual", 0), 0)

	def test_future_dated_payment_is_not_actual_or_settled_today(self):
		from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry

		order = fixtures.sales_order(self.context)
		invoice = fixtures.sales_invoice(order, 100)
		payment = get_payment_entry(
			"Sales Invoice", invoice.name, bank_account=self.context["bank"], party_amount=20
		)
		payment.posting_date = add_days(today(), 1)
		payment.reference_no = "Liquidity future test"
		payment.reference_date = payment.posting_date
		payment.insert()
		payment.submit()
		result = self.totals()
		self.assertEqual(result["total__forecast"], 100)
		self.assertEqual(result["total__plan"], 100)
		self.assertEqual(result.get("total__actual", 0), 0)

	def test_submitted_future_invoice_is_forecast_without_actual(self):
		from erpnext.selling.doctype.sales_order.sales_order import make_sales_invoice

		order = fixtures.sales_order(self.context)
		invoice = make_sales_invoice(order.name)
		invoice.set_posting_time = 1
		invoice.posting_date = add_days(today(), 1)
		invoice.due_date = add_days(today(), 7)
		invoice.set("payment_schedule", [])
		invoice.insert()
		invoice.submit()
		result = self.totals()
		self.assertEqual(result["total__forecast"], 100)
		self.assertEqual(result["total__plan"], 100)
		self.assertEqual(result.get("total__actual", 0), 0)

	def test_project_manager_api_details_options_and_export_share_scope(self):
		from frappe.share import add

		from dinext_projects_pl.liquidity.api import get_options

		order = fixtures.sales_order(self.context)
		invoice = fixtures.sales_invoice(order, 60)
		# A separate non-project order must not leak through the same company.
		hidden = fixtures.sales_order(self.context, amount=47, project=False)
		for doc in (order, invoice):
			add(doc.doctype, doc.name, user=self.manager, read=1)
		frappe.set_user(self.manager)
		try:
			result = calendar(dict(self.filters, grouping="party_only"))
			payload = get_calendar(dict(self.filters, grouping="party_only"), depth=6)
			self.assertIsNone(payload["balances"])
			self.assertEqual(payload["directions"]["in"]["total__forecast"], 60)
			self.assertNotIn(hidden.name, str(result.all_rows()))
			rows = get_cell_details(self.filters, "root", "total", "forecast")["rows"]
			self.assertTrue(all(r["analytic"]["project"] == self.context["project"] for r in rows))
			self.assertNotIn(hidden.name, str(get_options("party", fixtures.COMPANY, party_type="Customer")))
			workbook = load_workbook(BytesIO(make_workbook(result)))
			self.assertNotIn(hidden.name, str([[c.value for c in r] for s in workbook for r in s]))
		finally:
			frappe.set_user("Administrator")

	def test_portfolio_choices_respect_document_policy_after_access_changes(self):
		from dinext_projects_pl.liquidity.api import get_options

		original = frappe.get_hooks

		def changed_hooks(name=None, *args, **kwargs):
			result = original(name, *args, **kwargs)
			if name == "has_permission":
				result = dict(result)
				result["Project"] = [*result.get("Project", []), __name__ + ".deny_test_project"]
			return result

		frappe.set_user(self.manager)
		try:
			self.assertTrue(get_options("portfolio", fixtures.COMPANY))
			# An explicit DocShare still grants read access after a controller denial
			# in Frappe. Revoke that grant too when modelling loss of access.
			frappe.set_user("Administrator")
			frappe.share.remove("Project", self.context["project"], self.manager)
			frappe.set_user(self.manager)
			with patch("frappe.get_hooks", side_effect=changed_hooks):
				self.assertEqual(get_options("portfolio", fixtures.COMPANY), [])
				self.assertEqual(get_options("project", fixtures.COMPANY), [])
		finally:
			frappe.set_user("Administrator")

	def test_xlsx_escapes_formulas_preserves_money_and_full_tree(self):
		from dataclasses import replace

		order = fixtures.sales_order(self.context)
		fixtures.sales_invoice(order, 60)
		result = calendar(self.filters)
		flow = result.visible[0]
		flow = replace(flow, analytic=replace(flow.analytic, party="=SUM(1,2)"))
		from dinext_projects_pl.liquidity.presentation import Calendar

		result = Calendar(result.filters, result.snapshot, [flow])
		book = load_workbook(BytesIO(make_workbook(result)))
		self.assertTrue(any(c.value == "'=SUM(1,2)" for s in book for r in s for c in r))
		self.assertFalse(any(c.data_type == "f" for s in book for r in s for c in r))
		self.assertEqual(book.worksheets[1]["D2"].data_type, "n")

	def test_foreign_invoice_request_and_base_currency_payment(self):
		invoice = fixtures.purchase_invoice(self.context, currency="EUR", rate=2)
		request = fixtures.payment_request(invoice, 60, add_days(today(), 21))
		fixtures.pay(self.context, invoice, 20, request, bank_amount=40)
		result = get_calendar(self.filters)
		self.assertEqual(result["directions"]["out"]["total__forecast"], 160)
		self.assertEqual(result["directions"]["out"]["total__actual"], 40)
		rows = get_cell_details(self.filters, "root", "total", "forecast")["rows"]
		self.assertTrue(all(r["currency"] == "EUR" for r in rows))
		self.assertEqual(sum(r["original_amount"] for r in rows), 80)

	def test_invalid_legacy_exchange_rate_is_explicitly_incomplete(self):
		invoice = fixtures.purchase_invoice(self.context)
		invoice.db_set("conversion_rate", 0)
		result = get_calendar(self.filters)
		self.assertFalse(result["complete"])
		self.assertIsNone(result["balances"])
		self.assertTrue(any("exchange rate" in w for w in result["warnings"]))

	def test_explicit_payment_order_coverage_replaces_standalone_forecast(self):
		order = fixtures.sales_order(self.context)
		request = fixtures.payment_request(order, 100, add_days(today(), 14))
		first = fixtures.sales_invoice(order, 60)
		fixtures.payment_order(self.context, request, invoice=first, covered=60)
		rows = get_cell_details(self.filters, "root", "total", "forecast")["rows"]
		self.assertEqual(sum(r["amount"] for r in rows), 100)
		self.assertEqual(sum(r["amount"] for r in rows if r["source"]["doctype"] == "Payment Request"), 40)
		second = fixtures.sales_invoice(order, 40)
		fixtures.payment_order(self.context, request, invoice=second, covered=40)
		rows = get_cell_details(self.filters, "root", "total", "forecast")["rows"]
		self.assertEqual(sum(r["amount"] for r in rows), 100)
		self.assertTrue(all(r["source"]["doctype"] == "Sales Invoice" for r in rows))

	def test_explicit_coverage_after_consumed_request_advance(self):
		order = fixtures.sales_order(self.context)
		request = fixtures.payment_request(order, 100, add_days(today(), 14))
		fixtures.pay(self.context, order, 30, request)
		invoice = fixtures.sales_invoice(order, 60, advances=True)
		fixtures.payment_order(self.context, request, invoice=invoice, covered=60)
		result = self.totals()
		self.assertEqual(result["total__forecast"], 70)
		self.assertEqual(result["total__actual"], 30)
		self.assertEqual(result["total__plan"], 70)
		rows = get_cell_details(self.filters, "root", "total", "forecast")["rows"]
		self.assertEqual(sum(r["amount"] for r in rows if r["source"]["doctype"] == "Sales Invoice"), 30)
		self.assertEqual(sum(r["amount"] for r in rows if r["source"]["doctype"] == "Payment Request"), 40)

	def test_multiple_orders_contracts_and_single_split_payment(self):
		from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry
		from erpnext.selling.doctype.sales_order.sales_order import make_sales_invoice

		first = fixtures.sales_order(self.context, 60)
		second = fixtures.sales_order(self.context, 40)
		contracts = {fixtures.contract(first).name, fixtures.contract(second).name}
		invoice = make_sales_invoice(second.name, target_doc=make_sales_invoice(first.name))
		invoice.posting_date = today()
		invoice.due_date = add_days(today(), 7)
		invoice.set("payment_schedule", [])
		invoice.insert()
		invoice.submit()
		third = fixtures.sales_order(self.context, 50)
		other_invoice = fixtures.sales_invoice(third, 50)
		payment = get_payment_entry(
			"Sales Invoice", invoice.name, bank_account=self.context["bank"], party_amount=50
		)
		payment.references[0].allocated_amount = 30
		payment.append(
			"references",
			{
				"reference_doctype": "Sales Invoice",
				"reference_name": other_invoice.name,
				"total_amount": 50,
				"outstanding_amount": 50,
				"allocated_amount": 20,
			},
		)
		payment.reference_no = "Liquidity split payment"
		payment.reference_date = today()
		payment.insert()
		payment.submit()
		result = self.totals()
		self.assertEqual(
			(result["total__plan"], result["total__forecast"], result["total__actual"]), (100, 100, 50)
		)
		rows = get_cell_details(self.filters, "root", "total", "actual")["rows"]
		self.assertEqual(
			{r["analytic"]["order"]["name"]: r["amount"] for r in rows},
			{first.name: 18, second.name: 12, third.name: 20},
		)
		self.assertTrue(contracts.issubset({r["analytic"]["contract"] for r in rows}))

	def test_noncash_revaluation_changes_balance_not_turnover(self):
		expense = frappe.db.get_value("Company", fixtures.COMPANY, "default_expense_account")
		fixtures.journal(
			self.context,
			[
				{"account": self.context["bank"], "debit_in_account_currency": 7},
				{
					"account": expense,
					"credit_in_account_currency": 7,
					"cost_center": frappe.db.get_value("Company", fixtures.COMPANY, "cost_center"),
				},
			],
			voucher_type="Exchange Gain Or Loss",
		)
		result = get_calendar(self.filters)
		self.assertEqual(result["directions"]["in"].get("total__actual", 0), 0)
		self.assertEqual(result["directions"]["in"]["total__adjustment"], 7)
		self.assertEqual(result["balances"]["current_cash"], 7)

	def test_paid_request_status_without_cash_is_not_actual(self):
		order = fixtures.sales_order(self.context)
		request = fixtures.payment_request(order, 100)
		request.db_set("status", "Paid")
		result = self.totals()
		self.assertEqual(result["total__forecast"], 100)
		self.assertEqual(result.get("total__actual", 0), 0)

	def test_fixture_helpers_require_test_site_and_authorized_operator(self):
		from unittest.mock import patch

		from dinext_projects_pl.tests.liquidity_ui_helpers import guard

		with patch.dict(frappe.conf, {"allow_tests": False}):
			with self.assertRaises(frappe.PermissionError):
				guard()
		frappe.set_user(self.manager)
		try:
			with self.assertRaises(frappe.PermissionError):
				guard()
		finally:
			frappe.set_user("Administrator")
