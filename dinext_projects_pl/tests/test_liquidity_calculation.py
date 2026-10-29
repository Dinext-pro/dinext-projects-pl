import unittest
from dataclasses import replace
from datetime import date, timedelta
from decimal import Decimal as D

from dinext_projects_pl.liquidity.calculation import calculate
from dinext_projects_pl.liquidity.model import (
	Analytic,
	Filters,
	Flow,
	Link,
	Obligation,
	Request,
	Share,
	Snapshot,
	Stage,
	allocate,
	periods,
	remaining_stages,
)
from dinext_projects_pl.liquidity.presentation import Calendar, node_key

TODAY = date(2026, 10, 28)
INVOICE = Link("Sales Invoice", "INV-1")
ORDER = Link("Sales Order", "SO-1")
REQUEST = Link("Payment Request", "PR-1")
ANALYTIC = Analytic("Project", "Customer", "Customer", "Contract", ORDER, INVOICE)


def invoice(total=100, outstanding=100, due=TODAY, shares=None):
	return Obligation(
		INVOICE,
		"in",
		D(total),
		D(outstanding),
		"USD",
		D(1),
		[Stage("1", due, D(total))],
		shares or [Share("1", D(1), ANALYTIC)],
	)


def request(amount=100, due=TODAY, covered=None):
	return Request(
		REQUEST,
		"in",
		D(amount),
		D(amount),
		"USD",
		D(1),
		due,
		[Share("1", D(1), replace(ANALYTIC, invoice=None))],
		covered or {},
	)


def actual(amount=20, day=TODAY):
	return Flow(
		"GL-1", "actual", "in", D(amount), day, day, Link("Payment Entry", "PE-1"), ANALYTIC, "USD", D(amount)
	)


def filters(**kwargs):
	return Filters.parse(
		dict(
			company="Company",
			from_date=TODAY.isoformat(),
			to_date=(TODAY + timedelta(days=90)).isoformat(),
			**kwargs,
		),
		TODAY,
	)


class TestLiquidityCalculation(unittest.TestCase):
	def test_export_labels_follow_each_period_and_mixed_total(self):
		from io import BytesIO
		from unittest.mock import patch

		from openpyxl import load_workbook

		from dinext_projects_pl.liquidity.api import make_workbook

		snapshot = Snapshot(TODAY, "USD", invoices=[invoice(60, 40)], actual=[actual(20)])
		for metrics in (("actual", "forecast"), ("actual",), ("forecast",)):
			result = Calendar(
				replace(filters(), from_date=TODAY - timedelta(days=14), metrics=metrics),
				snapshot,
				calculate(snapshot),
			)
			with patch("dinext_projects_pl.liquidity.api._", lambda text: text):
				headers = next(iter(load_workbook(BytesIO(make_workbook(result))).active.values))[2:]
			if len(metrics) == 2:
				self.assertEqual(headers[0], "Without date / Forecast")
				self.assertTrue(any(label.endswith(" / Actual") for label in headers))
				self.assertTrue(any(label.endswith(" / Forecast") for label in headers))
				self.assertEqual(headers[-1], "Total / Cash movement")
			else:
				self.assertTrue(all(label.endswith(" / " + metrics[0].title()) for label in headers))

	def test_party_only_combines_projects_without_changing_flow_shares(self):
		flows = [
			actual("20.25"),
			replace(actual("30.75"), id="GL-2", analytic=replace(ANALYTIC, project="Other")),
		]
		snapshot = Snapshot(TODAY, "USD")
		for grouping in ("project_party", "party_project", "party_only"):
			result = Calendar(filters(grouping=grouping), snapshot, flows)
			self.assertEqual(result.details("root", "total", "actual")["total"], 51)
			if grouping == "party_only":
				rows = result.all_rows()
				self.assertNotIn("project", [r["level"] for r in rows])
				parties = [r for r in rows if r["level"] == "party"]
				self.assertEqual(len(parties), 1)
				self.assertEqual(result.details(parties[0]["id"], "total", "actual")["count"], 2)
				self.assertEqual(max(r["depth"] for r in rows), 5)

	def test_tree_depth_keeps_lazy_boundaries_and_all_terminal_metrics(self):
		flows = [
			actual(),
			replace(actual(80), id="plan", metric="plan"),
			replace(actual(40), id="forecast", metric="forecast"),
		]
		for grouping, maximum in (("project_party", 7), ("party_only", 6)):
			result = Calendar(filters(grouping=grouping), Snapshot(TODAY, "USD"), flows)

			def flatten(rows):
				return [item for row in rows for item in [row, *flatten(row.get("children", []))]]

			for depth in (1, 2, maximum):
				rows = flatten(result.tree(depth))
				self.assertEqual(max(row["depth"] for row in rows), depth - 1)
				self.assertEqual(rows[0]["total__plan"], 80)
				if depth < maximum:
					self.assertTrue(any(row["hasChildren"] and "children" not in row for row in rows))
				else:
					self.assertEqual({row["id"] for row in rows}, set(result.nodes))

	def test_combined_details_page_and_total_keep_each_source_once(self):
		s = Snapshot(
			TODAY,
			"USD",
			invoices=[invoice(60, 40)],
			actual=[actual(20), replace(actual(10, TODAY - timedelta(days=7)), id="GL-past")],
		)
		result = Calendar(replace(filters(), from_date=TODAY - timedelta(days=14)), s, calculate(s))
		first = result.details("root", "total", "combined", 0, 1)
		second = result.details("root", "total", "combined", 1, 1)
		self.assertEqual(first["total"], 50)
		self.assertEqual(first["count"], 2)
		self.assertNotEqual(first["rows"][0]["id"], second["rows"][0]["id"])
		self.assertEqual({first["rows"][0]["metric"], second["rows"][0]["metric"]}, {"actual", "forecast"})
		self.assertEqual(result.details("net", "total", "combined")["total"], 50)

	def test_combined_selection_across_day_week_and_month_boundaries(self):
		for grain in ("day", "week", "month"):
			with self.subTest(grain=grain):
				past_day = TODAY - timedelta(days=40)
				future_day = TODAY + timedelta(days=40)
				flows = [
					replace(actual(10, past_day), id="past"),
					replace(actual(20), id="current"),
					replace(actual(40), id="current-forecast", metric="forecast"),
					replace(actual(60, future_day), id="future-forecast", metric="forecast"),
				]
				result = Calendar(
					replace(filters(grain=grain), from_date=past_day), Snapshot(TODAY, "USD"), flows
				)
				for day, amount in ((past_day, 10), (TODAY, 40), (future_day, 60)):
					key = result.period_by_day[day]
					self.assertEqual(result.details("net", key, "combined")["total"], amount)
					self.assertEqual(result.display_value(result.values["root"], key, "combined"), amount)
				self.assertEqual(result.details("net", "total", "combined")["total"], 110)
				self.assertEqual(result.display_value(result.values["root"], "total", "combined"), 110)

	def test_combined_adjustments_details_match_only_completed_periods(self):
		for metric in ("transfer", "adjustment"):
			past = replace(actual(10, TODAY - timedelta(days=7)), id="past", metric=metric)
			current = replace(actual(20), metric=metric)
			s = Snapshot(TODAY, "USD")
			result = Calendar(replace(filters(), from_date=TODAY - timedelta(days=14)), s, [past, current])
			self.assertEqual(result.details("net", "total", "combined_" + metric)["total"], 10)
			self.assertEqual(
				result.details("net", result.period_by_day[TODAY], "combined_" + metric)["count"], 0
			)
			self.assertEqual(result.details("net", "total", metric)["total"], 30)

	def test_combined_export_has_plan_then_cash_with_undated_first(self):
		from io import BytesIO
		from unittest.mock import patch

		from openpyxl import load_workbook

		from dinext_projects_pl.liquidity.api import make_workbook

		s = Snapshot(TODAY, "USD", invoices=[invoice(60, 40)], actual=[actual(20)])
		result = Calendar(filters(), s, calculate(s))
		with patch("dinext_projects_pl.liquidity.api._", side_effect=lambda text: text):
			book = load_workbook(BytesIO(make_workbook(result)))
		rows = list(book.worksheets[0].values)
		self.assertEqual(rows[0][2:4], ("Without date / Plan", "Without date / Forecast"))
		self.assertEqual(rows[0][-2:], ("Total / Plan", "Total / Forecast"))
		self.assertEqual(rows[1][-1], 40)
		self.assertIn("44 · 2026", rows[0][4])
		components = list(book.worksheets[1].values)[1:]
		self.assertEqual({row[1] for row in components}, {"forecast"})
		self.assertEqual(sum(row[3] for row in components), 40)

	def test_current_negative_balance_is_not_hidden_by_same_day_receipt(self):
		s = Snapshot(TODAY, "USD", current_cash=D(-10), can_balance=True, invoices=[invoice(20, 20)])
		result = Calendar(filters(), s, calculate(s))
		self.assertEqual(result.balances["minimum"], -10)
		self.assertEqual(result.balances["first_gap_date"], TODAY.isoformat())

	def test_request_without_date_keeps_invoice_due_date(self):
		due = TODAY + timedelta(days=7)
		s = Snapshot(
			TODAY,
			"USD",
			invoices=[invoice(due=due)],
			requests=[replace(request(due=None, covered={INVOICE: D(100)}), invoice_only=True)],
		)
		flows = calculate(s)
		self.assertEqual(len(flows), 1)
		self.assertEqual(flows[0].date, due)
		self.assertEqual(flows[0].date_source, INVOICE)

	def test_future_request_preserves_contractual_overdue_flag(self):
		s = Snapshot(
			TODAY,
			"USD",
			invoices=[invoice(due=TODAY - timedelta(days=7))],
			requests=[
				replace(request(due=TODAY + timedelta(days=7), covered={INVOICE: D(100)}), invoice_only=True)
			],
		)
		flow = calculate(s)[0]
		self.assertFalse(flow.overdue)
		self.assertTrue(flow.contract_overdue)
		self.assertEqual(flow.date_source, REQUEST)

	def test_net_drilldown_has_signed_reconciled_total(self):
		s = Snapshot(TODAY, "USD", actual=[actual(30), replace(actual(10), id="GL-2", direction="out")])
		result = Calendar(filters(), s, calculate(s))
		detail = result.details("net", "total", "actual")
		self.assertTrue(detail["signed"])
		self.assertEqual(detail["total"], result.payload()["net"]["total__actual"])
		self.assertEqual(detail["total"], 20)

	def test_order_invoice_partial_payment(self):
		s = Snapshot(
			TODAY,
			"USD",
			orders=[replace(invoice(100, 80), source=ORDER)],
			invoices=[invoice(60, 40)],
			actual=[actual()],
		)
		flows = calculate(s)
		self.assertEqual(
			{m: sum(f.amount for f in flows if f.metric == m) for m in ("plan", "forecast", "actual")},
			{"plan": D(80), "forecast": D(40), "actual": D(20)},
		)

	def test_standalone_request_and_partial_payment(self):
		s = Snapshot(TODAY, "USD", requests=[request(70)], actual=[actual(30)])
		flows = calculate(s)
		self.assertEqual(sum(f.amount for f in flows if f.metric == "forecast"), 70)
		self.assertIsNone(next(f for f in flows if f.metric == "forecast").analytic.invoice)

	def test_partial_then_full_invoice_coverage(self):
		for covered in (60, 100):
			with self.subTest(covered=covered):
				s = Snapshot(
					TODAY,
					"USD",
					requests=[request(100, covered={INVOICE: D(covered)})],
					invoices=[invoice(covered, covered)],
				)
				flows = calculate(s)
				self.assertEqual(sum(f.amount for f in flows), 100)
				self.assertEqual(sum(f.amount for f in flows if f.analytic.invoice is None), 100 - covered)

	def test_request_refines_part_only(self):
		later = TODAY + timedelta(days=10)
		r = replace(request(40, later, {INVOICE: D(40)}), invoice_only=True)
		s = Snapshot(TODAY, "USD", invoices=[invoice(100, 80)], requests=[r])
		flows = calculate(s)
		self.assertEqual(
			{day: sum(f.amount for f in flows if f.date == day) for day in (TODAY, later)},
			{TODAY: D(40), later: D(40)},
		)

	def test_overlap_keeps_invoice_schedule(self):
		a = replace(request(80, TODAY + timedelta(days=2), {INVOICE: D(80)}), invoice_only=True)
		b = replace(a, source=Link("Payment Request", "PR-2"))
		s = Snapshot(TODAY, "USD", invoices=[invoice()], requests=[a, b])
		flows = calculate(s)
		self.assertEqual(sum(f.amount for f in flows), 100)
		self.assertTrue(all(f.date == TODAY for f in flows))
		self.assertFalse(s.complete)

	def test_overdue_priority_and_original_date(self):
		past, future = TODAY - timedelta(days=10), TODAY + timedelta(days=10)
		r = replace(request(60, future, {INVOICE: D(60)}), invoice_only=True)
		flows = calculate(Snapshot(TODAY, "USD", invoices=[invoice(due=past)], requests=[r]))
		self.assertEqual(sum(f.amount for f in flows if f.overdue), 40)
		self.assertEqual(sum(f.amount for f in flows if f.date == future), 60)
		self.assertTrue(all(f.original_date == past for f in flows))

	def test_no_date_outside_total_and_balance(self):
		s = Snapshot(TODAY, "USD", invoices=[invoice(due=None)], current_cash=D(50), can_balance=True)
		c = Calendar(filters(), s, calculate(s))
		p = c.payload()
		self.assertEqual(p["directions"]["in"]["undated__forecast"], 100)
		self.assertEqual(p["kpis"]["expected_receipts"], 0)
		self.assertEqual(p["balances"]["minimum"], 50)
		self.assertEqual(c.details("root", "total", "forecast")["count"], 0)

	def test_explicit_term_then_fifo_settlement(self):
		stages = [Stage("1", TODAY, D(40), "first"), Stage("2", TODAY + timedelta(days=1), D(60), "second")]
		result = remaining_stages(stages, D(100), D(70), {"second": D(20)})
		self.assertEqual([s.amount for s in result], [D(30), D(40)])
		self.assertEqual(sum(s.amount for s in result), D(70))

	def test_paid_and_overpaid_do_not_create_negative_forecast(self):
		for value in (0, -20):
			s = Snapshot(TODAY, "USD", invoices=[invoice(outstanding=value)], actual=[actual(100)])
			self.assertEqual([f.metric for f in calculate(s)], ["actual"])

	def test_stable_rounding_in_multiple_precisions(self):
		for precision in (0, 2, 3):
			parts = allocate(D(1), [D(1)] * 3, precision)
			self.assertEqual(sum(parts), D(1))
			self.assertEqual(parts, allocate(D(1), [D(1)] * 3, precision))
		self.assertEqual(allocate(D("0.02"), [D(1)] * 3), [D(0), D(0), D("0.02")])

	def test_duplicate_id_rejected(self):
		with self.assertRaises(ValueError):
			calculate(Snapshot(TODAY, "USD", actual=[actual(), actual()]))

	def test_future_posting_not_actual(self):
		self.assertEqual(
			calculate(Snapshot(TODAY, "USD", actual=[actual(day=TODAY + timedelta(days=1))])), []
		)

	def test_future_window_includes_intermediate_forecast(self):
		r = replace(request(80, TODAY + timedelta(days=2)), direction="out")
		s = Snapshot(TODAY, "USD", requests=[r], current_cash=D(100), can_balance=True)
		f = replace(filters(), from_date=TODAY + timedelta(days=10))
		c = Calendar(f, s, calculate(s))
		self.assertEqual(c.payload()["balances"]["periods"][0]["projected_opening"], 20)
		self.assertEqual(c.payload()["kpis"]["expected_payments"], 0)

	def test_daily_gap_not_hidden_by_weekly_closing(self):
		a = replace(request(150), direction="out")
		b = replace(request(200, TODAY + timedelta(days=1)), source=Link("Payment Request", "PR-2"))
		s = Snapshot(TODAY, "USD", requests=[a, b], current_cash=D(100), can_balance=True)
		p = Calendar(filters(), s, calculate(s)).payload()
		self.assertEqual(p["balances"]["minimum"], -50)
		self.assertEqual(p["balances"]["first_gap_date"], TODAY.isoformat())
		self.assertEqual(p["balances"]["periods"][0]["projected_closing"], 150)

	def test_accounts_hide_projected_only(self):
		s = Snapshot(TODAY, "USD", current_cash=D(100), can_balance=True)
		p = Calendar(filters(accounts=["Bank"]), s, []).payload()
		self.assertEqual(p["balances"]["current_cash"], 100)
		self.assertIsNone(p["balances"]["minimum"])
		self.assertTrue(all(r["projected_closing"] is None for r in p["balances"]["periods"]))

	def test_analytical_filter_hides_balances(self):
		s = Snapshot(TODAY, "USD", can_balance=True)
		self.assertIsNone(Calendar(filters(projects=["Project"]), s, []).payload()["balances"])

	def test_empty_company_retains_balance(self):
		s = Snapshot(TODAY, "USD", opening=D(80), current_cash=D(80), can_balance=True)
		p = Calendar(filters(), s, []).payload()
		self.assertEqual(p["rows"], [])
		self.assertEqual(p["balances"]["periods"][0]["actual_closing"], 80)

	def test_tree_reordering_and_drilldown_conserve_amounts(self):
		shares = [
			Share("a", D(1), ANALYTIC),
			Share("b", D(2), replace(ANALYTIC, project="Проект / <x>", contract="=1+1")),
		]
		s = Snapshot(TODAY, "USD", invoices=[invoice(shares=shares)])
		flows = calculate(s)
		for grouping in ("project_party", "party_project"):
			c = Calendar(filters(grouping=grouping), s, flows)
			self.assertEqual(c.children()[0]["total__forecast"], 100)
			self.assertEqual(sum(r["total__forecast"] for r in c.all_rows() if r["level"] == "contract"), 100)
			self.assertEqual(c.details("root", "total", "forecast", 0, 1)["total"], 100)
			self.assertEqual(c.details("root", "total", "forecast", 0, 1)["count"], 2)

	def test_own_metrics_stop_at_their_source_level(self):
		order = replace(invoice(), source=ORDER, shares=[Share("1", D(1), replace(ANALYTIC, invoice=None))])
		snapshot = Snapshot(TODAY, "USD", orders=[order], invoices=[invoice()])
		calendar = Calendar(filters(), snapshot, calculate(snapshot))
		rows = calendar.all_rows()
		self.assertFalse(any(r["level"] == "payment" for r in rows))
		invoice_rows = [r for r in rows if r["level"] == "invoice"]
		self.assertEqual(len(invoice_rows), 1)
		self.assertEqual(invoice_rows[0]["total__forecast"], 100)
		self.assertEqual(invoice_rows[0].get("total__plan", 0), 0)
		self.assertTrue(next(r for r in rows if r["level"] == "order")["hasChildren"])

	def test_unlinked_request_in_both_directions(self):
		for direction in ("in", "out"):
			r = replace(request(), direction=direction, shares=[Share("unlinked", D(1), Analytic())])
			s = Snapshot(TODAY, "USD", requests=[r], can_balance=True)
			c = Calendar(filters(), s, calculate(s))
			self.assertEqual(c.children()[0]["total__forecast"], 100)
			self.assertEqual(
				c.balances["periods"][0]["projected_closing"], 100 if direction == "in" else -100
			)
			self.assertTrue(any(row["label"] == "Without invoice" for row in c.all_rows()))
			self.assertTrue(any(row["label"] == "Without order" for row in c.all_rows()))

	def test_invalid_cells_and_pagination(self):
		c = Calendar(filters(), Snapshot(TODAY, "USD"), [])
		for key, period, metric, start, size in (
			("forged", "total", "actual", 0, 50),
			("root", "bad", "actual", 0, 50),
			("root", "total", "sum", 0, 50),
			("root", "total", "actual", -1, 50),
			("root", "total", "actual", 0, 201),
			("root", "total", "actual", 0.5, 50),
		):
			with self.assertRaises(ValueError):
				c.details(key, period, metric, start, size)

	def test_period_boundaries_leap_day_and_year(self):
		for grain in ("day", "week", "month"):
			f = replace(filters(grain=grain), from_date=date(2023, 12, 30), to_date=date(2024, 3, 2))
			ps = periods(f)
			self.assertEqual(ps[0].start, f.from_date)
			self.assertEqual(ps[-1].end, f.to_date)
			self.assertEqual(sum((p.end - p.start).days + 1 for p in ps), (f.to_date - f.from_date).days + 1)
			self.assertEqual(sum(p.start <= date(2024, 2, 29) <= p.end for p in ps), 1)

	def test_filter_validation(self):
		for raw in (
			{},
			{"company": "C", "accounts": "Bank"},
			{"company": "C", "metrics": ["sql"]},
			{"company": "C", "from_date": "2026-02-30"},
			{"company": "C", "from_date": "2026-11-30", "to_date": "2026-01-01"},
			{"company": "C", "ignore_permissions": True},
		):
			with self.assertRaises(ValueError):
				Filters.parse(raw, TODAY)

	def test_default_thirteen_weeks(self):
		f = Filters.parse({"company": "Company"}, TODAY)
		self.assertEqual(f.from_date.weekday(), 0)
		self.assertEqual(f.to_date.weekday(), 6)
		self.assertEqual(len(periods(f)), 13)

	def test_zero_rows_are_explicit_and_do_not_add_amounts(self):
		s = Snapshot(TODAY, "USD", invoices=[invoice(outstanding=0)])
		self.assertEqual(calculate(s), [])
		c = Calendar(filters(show_zero=True), s, calculate(s, include_zero=True))
		self.assertTrue(c.children())
		self.assertEqual(c.payload()["directions"]["in"]["total__forecast"], 0)

	def test_balance_drilldown_reconciles_to_anchor(self):
		s = Snapshot(
			TODAY,
			"USD",
			invoices=[invoice()],
			actual=[actual(20)],
			current_cash=D(70),
			opening=D(50),
			cash_daily={TODAY: D(20)},
			can_balance=True,
		)
		c = Calendar(filters(), s, calculate(s))
		for metric in ("actual_closing", "projected_closing"):
			d = c.details("root", c.periods[0].key, metric)
			self.assertEqual(d["total"], d["anchor"] + d["movement"])

	def test_split_document_details_do_not_disclose_sibling_amount(self):
		shares = [Share("a", D(1), ANALYTIC), Share("b", D(1), replace(ANALYTIC, project="Hidden"))]
		flows = calculate(Snapshot(TODAY, "USD", invoices=[invoice(100, 60, shares=shares)]))
		self.assertEqual(flows[0].amount, D(30))
		self.assertEqual(flows[0].paid, D(20))
		self.assertEqual(flows[0].outstanding, D(30))


if __name__ == "__main__":
	unittest.main()
